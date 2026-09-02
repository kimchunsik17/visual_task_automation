"""documents/format_import.py — 서식 파일(.hwpx/.docx) → FormatSpec 초안 (계획 보류 항목 '역변환').

포맷 스튜디오의 세 시작점(프리셋·AI 생성·빈 포맷)에 네 번째 — **가진 파일에서 시작** — 를 더한다.
파일의 문단·표 구조와 `{{자리표시자}}` 를 결정적으로 추출해 FormatSpec 초안을 만든다. 저장이
아니라 스튜디오 편집기에 로드되는 초안이라는 계약은 AI 생성(`format_studio.py`)과 같다.

원칙 두 가지:
- **파일은 읽기만 한다.** 업로드 저장소·artifact 로 남기지 않는다(호출자가 임시 파일로 준다).
- **결정적 추출이 정본이다.** 자리표시자가 없는 일반 서식의 "빈칸 찾기"는 이 초안을 근거로 한
  AI 다듬기(format_studio.refine_imported_spec)가 하고, 실패해도 이 초안은 항상 남는다.

한글 자리표시자(`{{작성자}}`)는 FormatSpec 필드 이름 규칙(영문·숫자·밑줄)에 맞지 않으므로
안전한 이름(field1…)으로 바꾸고 원문을 라벨로 보존한다 — 골격의 참조도 함께 고쳐 쓴다.
"""

from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Tuple
from xml.etree import ElementTree as ET

from . import hwpx
from .format_spec import FormatSpecError, MAX_FIELDS, validate_format_spec
from .hwpx import placeholders as hwpx_placeholders
from .hwpx import xmlio

IMPORT_EXTENSIONS = (".hwpx", ".docx")

# 추출 단계의 자리표시자 인식은 hwpx 채우기와 같은 규칙을 쓴다 — 이름에 공백만 없으면
# 한글도 잡는다. (FormatSpec 의 참조 규칙은 ASCII 라서 아래에서 이름을 바꾼다.)
_PLACEHOLDER_RE = hwpx_placeholders.PLACEHOLDER_RE
_SAFE_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,63}$")

_DOCX_HEADING_STYLE_RE = re.compile(r"^(?:Heading|heading|제목)\s*([1-9])")


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _paragraph_text(paragraph: ET.Element) -> str:
    """hwpx 문단의 논리 문자열 — 여러 run 으로 쪼개진 텍스트를 placeholders 와 같은 규칙으로 잇는다."""
    logical, _slots = hwpx_placeholders._slots_for(paragraph)
    # 제어 요소(탭 등) 자리의 경계 문자는 사람이 읽는 공백으로 되돌린다.
    return logical.replace(hwpx_placeholders._BARRIER, " ").strip()


def _hwpx_cell_text(cell: ET.Element) -> str:
    """표 셀 텍스트 — 셀 안의 (중첩 포함) 문단들을 줄바꿈으로 잇는다."""
    parts = [_paragraph_text(p) for p in cell.iter() if _local(p.tag) == "p"]
    return "\n".join(t for t in parts if t)


def _hwpx_table_block(table: ET.Element) -> Dict[str, Any] | None:
    grid: List[List[str]] = []
    for row in table:
        if _local(row.tag) != "tr":
            continue
        cells = [_hwpx_cell_text(tc) for tc in row if _local(tc.tag) == "tc"]
        if cells:
            grid.append(cells)
    return _table_block_from_grid(grid)


def _table_block_from_grid(grid: List[List[str]]) -> Dict[str, Any] | None:
    """행렬 → table 블록. 첫 행을 열 이름으로 본다(가져오기 관례 — 스튜디오에서 고칠 수 있다).

    병합 셀 등으로 행마다 칸 수가 다르면 렌더러 규칙(칸 수 = columns)에 맞게 빈 칸으로
    채운다 — 조용히 셀을 버리는 것보다 낫고, 무엇이 채워졌는지 미리보기에서 바로 보인다.
    """
    grid = [row for row in grid if any(cell.strip() for cell in row)]
    if not grid:
        return None
    width = max(len(row) for row in grid)
    padded = [row + [""] * (width - len(row)) for row in grid]
    return {"type": "table", "columns": padded[0], "rows": padded[1:]}


def extract_blocks_hwpx(path: str) -> List[Dict[str, Any]]:
    """섹션 XML 을 문서 순서로 걷는다: 최상위 문단의 텍스트·표·쪽나눔만 본다(셀 안 문단은 표가 담당)."""
    package = hwpx.HwpxPackage.open(path)   # zip bomb·XML 폭탄 검사 포함
    blocks: List[Dict[str, Any]] = []
    for name in package.section_names():
        root, _header = xmlio.parse(package.read(name), name=name)
        for paragraph in root:
            if _local(paragraph.tag) != "p":
                continue
            if paragraph.get("pageBreak") == "1":
                blocks.append({"type": "page_break"})
            text = _paragraph_text(paragraph)
            if text:
                blocks.append({"type": "paragraph", "text": text})
            for run in paragraph:
                if _local(run.tag) != "run":
                    continue
                for child in run:
                    if _local(child.tag) == "tbl":
                        table = _hwpx_table_block(child)
                        if table:
                            blocks.append(table)
    return blocks


def extract_blocks_docx(path: str) -> List[Dict[str, Any]]:
    """python-docx 로 본문을 문서 순서(iter_inner_content)로 걷는다. 제목 스타일은 heading 으로 살린다."""
    try:
        from docx import Document
    except ImportError as exc:  # requirements 에 있으므로 배포 환경에서는 나오지 않는다
        raise FormatSpecError(f"docx 를 읽을 수 없습니다(python-docx 미설치): {exc}") from None

    try:
        document = Document(path)
    except Exception as exc:
        raise FormatSpecError(f"Word 문서를 열지 못했습니다: {exc}") from None

    blocks: List[Dict[str, Any]] = []
    for item in document.iter_inner_content():
        kind = type(item).__name__
        if kind == "Paragraph":
            text = (item.text or "").strip()
            if not text:
                continue
            style_name = getattr(getattr(item, "style", None), "name", "") or ""
            match = _DOCX_HEADING_STYLE_RE.match(style_name)
            if match:
                blocks.append({"type": "heading", "level": min(3, int(match.group(1))), "text": text})
            else:
                blocks.append({"type": "paragraph", "text": text})
        elif kind == "Table":
            grid = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            table = _table_block_from_grid(grid)
            if table:
                blocks.append(table)
    return blocks


def _iter_block_texts(blocks: List[Dict[str, Any]]):
    """블록 안에서 자리표시자가 살 수 있는 모든 문자열 자리 — (컨테이너, 키/인덱스) 로 돌려준다."""
    for block in blocks:
        if isinstance(block.get("text"), str):
            yield block, "text"
        for row in ([block.get("columns")] if block.get("columns") else []) + list(block.get("rows") or []):
            for index in range(len(row)):
                if isinstance(row[index], str):
                    yield row, index


def _collect_and_rename_placeholders(blocks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """{{이름}} 들을 fields 로 선언하고, 규칙(ASCII) 밖 이름은 골격까지 함께 바꿔 쓴다."""
    names: List[str] = []
    for container, key in _iter_block_texts(blocks):
        for match in _PLACEHOLDER_RE.finditer(container[key]):
            if match.group(1) not in names:
                names.append(match.group(1))

    if len(names) > MAX_FIELDS:
        raise FormatSpecError(
            f"문서의 자리표시자가 너무 많습니다({len(names)}개, 상한 {MAX_FIELDS}개).")

    rename: Dict[str, str] = {}
    used: set = set()
    counter = 0
    for original in names:
        if _SAFE_NAME_RE.match(original) and original not in used:
            safe = original
        else:
            counter += 1
            while f"field{counter}" in used or f"field{counter}" in names:
                counter += 1
            safe = f"field{counter}"
        rename[original] = safe
        used.add(safe)

    changed = {o: s for o, s in rename.items() if o != s}
    if changed:
        def _substitute(text: str) -> str:
            return _PLACEHOLDER_RE.sub(
                lambda m: "{{" + rename.get(m.group(1), m.group(1)) + "}}", text)
        for container, key in _iter_block_texts(blocks):
            container[key] = _substitute(container[key])

    return [{"name": rename[original], "label": original, "kind": "text", "required": False}
            for original in names]


def spec_from_file(path: str, *, original_name: str = "") -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """파일 → (검증 통과한 FormatSpec 초안, 추출 요약). 실패는 FormatSpecError 로 올린다."""
    extension = os.path.splitext(original_name or path)[1].lower()
    if extension not in IMPORT_EXTENSIONS:
        raise FormatSpecError(
            f"지원하지 않는 파일 형식입니다: {extension or '(없음)'} — "
            f"{', '.join(IMPORT_EXTENSIONS)} 만 가져올 수 있습니다.")

    try:
        blocks = (extract_blocks_hwpx(path) if extension == ".hwpx"
                  else extract_blocks_docx(path))
    except hwpx.PackageRejected as exc:
        raise FormatSpecError(f"HWPX 문서를 열지 못했습니다: {exc}") from None

    if not any(block["type"] in ("paragraph", "heading", "table") for block in blocks):
        raise FormatSpecError("문서에서 가져올 내용(문단·표)을 찾지 못했습니다.")

    fields = _collect_and_rename_placeholders(blocks)

    stem = os.path.splitext(os.path.basename(original_name or path))[0].strip() or "가져온 포맷"
    default_output = "hwpx" if extension == ".hwpx" else "docx"
    spec = validate_format_spec({
        "version": 1,
        "layout": "document",
        "name": stem,
        "output": {"default": default_output, "allowed": ["hwpx", "docx", "pdf", "xlsx"]},
        "fields": fields,
        "blocks": blocks,
    })
    info = {
        "placeholders": [f["label"] for f in fields],
        "blocks": len(blocks),
        "tables": sum(1 for b in blocks if b["type"] == "table"),
    }
    return spec, info
